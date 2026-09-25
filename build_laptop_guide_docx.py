import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_laptop_guide_docx():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Document Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("LiveGuard-EHMS")
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Immediate Development Guide: Laptop + Arduino Sensor Prototype\n(No Raspberry Pi Required)")
    sub_run.font.size = Pt(14)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_custom_heading(text, level=1):
        h = doc.add_heading(text, level=level)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        for r in h.runs:
            r.font.name = 'Calibri'
            if level == 1:
                r.font.size = Pt(15)
                r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
                r.font.bold = True
            elif level == 2:
                r.font.size = Pt(12)
                r.font.color.rgb = RGBColor(0x2E, 0x6B, 0x9E)
                r.font.bold = True
        return h

    # Section 1: Executive Overview
    add_custom_heading("1. Overview: How Your Laptop Acts as the Edge Device")
    doc.add_paragraph(
        "While awaiting your Raspberry Pi, your Laptop functions as the complete Edge Intelligence Node. "
        "The Arduino acts strictly as a high-speed analog acquisition bridge, capturing real-time ECG signals "
        "at 360 Hz and streaming them over USB Serial into the Python pipeline running on your PC."
    )
    doc.add_paragraph(
        "The Laptop executes:\n"
        "• Digital Signal Filtering (0.5–40 Hz Butterworth Bandpass + 50 Hz Notch filter)\n"
        "• Real-Time Pan-Tompkins QRS / R-Peak Detection\n"
        "• 180-Sample Beat Normalization & Windowing\n"
        "• PyTorch Stage 1 CNN AI Inference (< 3ms per heartbeat)\n"
        "• WebSocket Live Telemetry Server (for browser dashboard display)"
    )

    # Section 2: Hardware Wiring
    add_custom_heading("2. Step 1: Wire the AD8232 to Your Arduino")
    doc.add_paragraph(
        "Use jumper wires to connect the AD8232 ECG sensor to your Arduino Uno or Nano as follows:"
    )

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ["AD8232 Sensor Pin", "Arduino Pin", "Wire Color", "Description"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8EEF5"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)

    wiring = [
        ["3.3V", "3.3V", "Red", "Power supply (Do NOT connect to 5V)"],
        ["GND", "GND", "Black", "Ground reference"],
        ["OUTPUT", "Pin A0", "Yellow", "Analog ECG voltage signal (0-1023)"],
        ["LO+", "Digital Pin 2", "Blue", "Leads-Off Detection Positive"],
        ["LO-", "Digital Pin 3", "Green", "Leads-Off Detection Negative"],
        ["SDN", "Unconnected", "—", "Leave open / unconnected"]
    ]

    for row in wiring:
        r_cells = table.add_row().cells
        for i, val in enumerate(row):
            r_cells[i].text = val
            r_cells[i].paragraphs[0].runs[0].font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 3: Electrode Placement
    add_custom_heading("3. Step 2: Stick the 3 Electrodes on Your Body")
    doc.add_paragraph(
        "Connect the 3.5mm electrode cable into the AD8232 jack and attach 3 disposable gel pads:\n"
        "1. RED (RA - Right Arm): Stick below your right collarbone (clavicle) or on right inner wrist.\n"
        "2. YELLOW (LA - Left Arm): Stick below your left collarbone (clavicle) or on left inner wrist.\n"
        "3. GREEN (RL - Right Leg): Stick on your lower right abdomen / ribcage (Ground reference electrode)."
    )
    doc.add_paragraph(
        "Important: Clean skin with an alcohol wipe to remove natural oils for clear waveforms with no drift."
    )

    # Section 4: Flashing Firmware
    add_custom_heading("4. Step 3: Upload Firmware to Arduino")
    doc.add_paragraph(
        "1. Open the Arduino IDE on your laptop.\n"
        "2. Open file: c:\\LiveGuard\\edge_system\\arduino_bridge\\liveguard_sensor_bridge.ino\n"
        "3. Plug the Arduino into your laptop via USB cable.\n"
        "4. Go to Tools -> Board -> Select 'Arduino Uno' (or Nano).\n"
        "5. Go to Tools -> Port -> Select your COM port (e.g., COM3, COM4).\n"
        "6. Click the Upload button (Arrow icon).\n"
        "7. Optional Test: Open Tools -> Serial Plotter, set baud rate to 115200 to see live heartbeats!"
    )

    # Section 5: Python Setup & Running
    add_custom_heading("5. Step 4: Run the LiveGuard Edge AI on Your Laptop")
    doc.add_paragraph(
        "1. Open PowerShell or Command Prompt in c:\\LiveGuard\n"
        "2. Install required Python packages:\n"
        "   pip install pyserial websockets torch scipy numpy scikit-learn\n\n"
        "3. Make sure to CLOSE the Arduino Serial Monitor / Serial Plotter so the COM port is free.\n\n"
        "4. Start the Edge AI engine with live sensors:\n"
        "   python -m edge_system.run_edge --source ARDUINO --port COM3\n\n"
        "5. Or test in Mock Mode without any wires connected:\n"
        "   python -m edge_system.run_edge --source MOCK"
    )

    # Section 6: Interpreting Output
    add_custom_heading("6. Understanding Terminal Outputs")
    doc.add_paragraph(
        "When the pipeline is running, it segments every heartbeat and displays:\n"
        "• Beat Number & Real-time Heart Rate (BPM)\n"
        "• AI Diagnosis: [NORMAL BEAT] (Conf: 98.2%) or [ABNORMAL/ARRHYTHMIA] (Prob: 0.88)\n"
        "• Total Arrhythmia Alerts count\n"
        "• WebSocket server status streaming at ws://0.0.0.0:8765 for the web dashboard."
    )

    # Section 7: Troubleshooting
    add_custom_heading("7. Troubleshooting & Common Fixes")
    doc.add_paragraph(
        "• 'PermissionError / Access Denied on COM port':\n"
        "   Fix: The Arduino Serial Monitor is open in Arduino IDE. Close it so Python can access the port.\n\n"
        "• 'Leads-off detected warning':\n"
        "   Fix: One or more electrode pads came loose from your skin. Re-press the gel pads.\n\n"
        "• 'Signal looks like a flatline':\n"
        "   Fix: Ensure AD8232 is powered by 3.3V, GND is shared, and OUTPUT is connected to Analog Pin A0."
    )

    output_path = "c:\\LiveGuard\\LiveGuard_Laptop_Arduino_Quickstart.docx"
    doc.save(output_path)
    print(f"Successfully generated Word document at: {output_path}")

if __name__ == "__main__":
    create_laptop_guide_docx()
